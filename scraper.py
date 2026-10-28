import argparse
import csv
import json
import logging
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

BASE_URL = "https://books.toscrape.com/catalogue/page-1.html"
RATINGS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}

log = logging.getLogger("scraper")


def make_session():
    s = requests.Session()
    s.headers["User-Agent"] = "Mozilla/5.0 (compatible; bookstore-scraper/1.0)"
    retry = Retry(total=4, backoff_factor=1, status_forcelist=[429, 500, 502, 503, 504])
    s.mount("https://", HTTPAdapter(max_retries=retry, pool_maxsize=16))
    s.mount("http://", HTTPAdapter(max_retries=retry, pool_maxsize=16))
    return s


def parse_price(text):
    m = re.search(r"[\d.]+", text or "")
    return float(m.group()) if m else None


def parse_listing(html, page_url):
    soup = BeautifulSoup(html, "html.parser")
    books = []
    for card in soup.select("article.product_pod"):
        link = card.select_one("h3 a")
        rating = card.select_one("p.star-rating")
        books.append({
            "title": link["title"],
            "price": parse_price(card.select_one(".price_color").text),
            "rating": RATINGS.get(rating["class"][1]) if rating else None,
            "in_stock": "In stock" in card.select_one(".availability").text,
            "url": urljoin(page_url, link["href"]),
        })

    next_link = soup.select_one("li.next a")
    next_url = urljoin(page_url, next_link["href"]) if next_link else None
    return books, next_url


def parse_book(html):
    soup = BeautifulSoup(html, "html.parser")
    table = {row.th.text.strip(): row.td.text.strip() for row in soup.select("table.table tr")}

    desc = soup.select_one("#product_description + p")
    crumbs = soup.select("ul.breadcrumb li a")
    stock = re.search(r"\((\d+) available\)", table.get("Availability", ""))

    return {
        "upc": table.get("UPC"),
        "category": crumbs[-1].text.strip() if len(crumbs) > 2 else None,
        "stock_count": int(stock.group(1)) if stock else 0,
        "reviews": int(table.get("Number of reviews", 0)),
        "description": desc.text.strip().removesuffix("...more").strip() if desc else "",
    }


def fetch(session, url, delay):
    resp = session.get(url, timeout=20)
    resp.raise_for_status()
    if delay:
        time.sleep(delay)
    # the site doesn't send a charset, so let bs4 sniff the bytes instead of resp.text
    return resp.content


def crawl(start_url, max_pages, details, workers, delay):
    session = make_session()
    url, page, books = start_url, 0, []

    while url and (max_pages is None or page < max_pages):
        page += 1
        log.info("page %d: %s", page, url)
        found, url = parse_listing(fetch(session, url, delay), url)
        books.extend(found)

    if details and books:
        log.info("fetching details for %d books (%d workers)", len(books), workers)

        def enrich(book):
            try:
                book.update(parse_book(fetch(session, book["url"], delay)))
            except requests.RequestException as e:
                log.warning("skipped details for %s: %s", book["url"], e)
            return book

        with ThreadPoolExecutor(max_workers=workers) as pool:
            books = list(pool.map(enrich, books))

    return books


def save(books, path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(dict.fromkeys(k for b in books for k in b))

    if path.suffix == ".json":
        path.write_text(json.dumps(books, indent=2, ensure_ascii=False), encoding="utf-8")
    elif path.suffix == ".xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Books"
        ws.append(fields)
        for b in books:
            ws.append([b.get(f) for f in fields])
            # scraped text starting with "=" would otherwise be saved as a live formula
            for cell in ws[ws.max_row]:
                if cell.data_type == "f":
                    cell.data_type = "s"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        wb.save(path)
    else:
        # utf-8-sig so Excel opens it without mangling characters
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=fields)
            w.writeheader()
            # Excel runs cells starting with = + - @ as formulas, so quote scraped text that does
            w.writerows({k: "'" + v if isinstance(v, str) and v[:1] in "=+-@" else v
                         for k, v in b.items()} for b in books)


def main(argv=None):
    p = argparse.ArgumentParser(description="Scrape the books.toscrape.com catalogue.")
    p.add_argument("-o", "--out", default="output/books.csv", help=".csv, .json or .xlsx")
    p.add_argument("-p", "--pages", type=int, help="stop after N listing pages (default: all)")
    p.add_argument("-d", "--details", action="store_true", help="also visit each book page")
    p.add_argument("--min-rating", type=int, default=0)
    p.add_argument("--max-price", type=float)
    p.add_argument("--workers", type=int, default=8)
    p.add_argument("--delay", type=float, default=0.2, help="seconds to wait after each request")
    p.add_argument("--url", default=BASE_URL, help="listing page to start from (e.g. a category)")
    args = p.parse_args(argv)

    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(message)s", datefmt="%H:%M:%S")

    started = time.time()
    try:
        books = crawl(args.url, args.pages, args.details, args.workers, args.delay)
    except requests.RequestException as e:
        log.error("crawl failed: %s", e)
        return 1

    books = [
        b for b in books
        if (b["rating"] or 0) >= args.min_rating
        and (args.max_price is None or b["price"] <= args.max_price)
    ]
    save(books, args.out)
    log.info("saved %d books to %s in %.1fs", len(books), args.out, time.time() - started)
    return 0


if __name__ == "__main__":
    sys.exit(main())
